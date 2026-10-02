# -*- coding: utf-8 -*-
import datetime as dt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = r"D:\SJH_Data\PersonalVault\01_Projects\API-인가-취약점-진단\4_회의\CCIT WBS.xlsx"
FONT = "맑은 고딕"
D = "%Y-%m-%d"

# (구분, 작업, 기능 상세, 담당자, 완료여부, 시작, 종료)
rows = [
    ("기획 및 설계", "주제 선정", "주제 확정 및 근거 수집 (API 인가 취약점 진단)", "전체", "완료", "2026-09-17", "2026-09-18"),
    ("기획 및 설계", "주제 선정", "주제 교체 제안서 작성", "심재훈", "완료", "2026-09-18", "2026-09-18"),
    ("기획 및 설계", "주제 선정", "교수·멘토 설득", "심재훈", "진행중", "2026-09-20", "2026-09-21"),
    ("기획 및 설계", "착수 검증", "crAPI 벤치 구동 + BOLA 재현", "심재훈", "완료", "2026-09-18", "2026-09-18"),
    ("기획 및 설계", "착수 검증", "샘플 3~4건 확대, 영향도 점수 차이 실증", "심재훈", "진행중", "2026-09-20", "2026-09-21"),
    ("기획 및 설계", "착수 검증", "mitmproxy + Newman 트래픽 20~30건 수집", "곽민경", "진행전", "2026-09-21", "2026-09-22"),
    ("기획 및 설계", "착수 검증", "LLM 판정 정확도 측정 (목표 70%)", "심재훈", "진행전", "2026-09-21", "2026-09-22"),
    ("기획 및 설계", "기술 스택 확정", "mitmproxy 리버스 프록시 게이트웨이 구성 확정", "전체", "완료", "2026-09-19", "2026-09-19"),
    ("기획 및 설계", "WBS 제작", "WBS 제작 및 역할·범위 확정", "전체", "진행중", "2026-09-20", "2026-09-21"),

    ("메인 기능", "1단계 엔드포인트 복원", "mitmproxy 애드온 수집 파이프라인", "곽민경", "진행전", "2026-09-24", "2026-10-07"),
    ("메인 기능", "1단계 엔드포인트 복원", "경로 변수 일반화, 엔드포인트 템플릿 복원", "곽민경", "진행전", "2026-10-01", "2026-10-14"),
    ("메인 기능", "1단계 엔드포인트 복원", "인증 헤더·쿠키 식별, 세션 컨텍스트 분리", "곽민경", "진행전", "2026-10-08", "2026-10-21"),
    ("메인 기능", "1단계 엔드포인트 복원", "엔드포인트 인벤토리 스키마 확정", "곽민경", "진행전", "2026-10-15", "2026-10-21"),
    ("메인 기능", "2단계 인가 취약점 탐지", "정적 프리필터 (LLM 호출 수 억제)", "심재훈", "진행전", "2026-09-24", "2026-10-07"),
    ("메인 기능", "2단계 인가 취약점 탐지", "세션 교차 재생 엔진 (A 리소스를 B 토큰으로 호출)", "심재훈", "진행전", "2026-10-01", "2026-10-21"),
    ("메인 기능", "2단계 인가 취약점 탐지", "LLM 판정 프롬프트·출력 스키마 확정", "심재훈", "진행전", "2026-10-08", "2026-10-28"),
    ("메인 기능", "2단계 인가 취약점 탐지", "BOLA·BFLA 판정 결과 구조화", "심재훈", "진행전", "2026-10-22", "2026-11-04"),
    ("메인 기능", "3단계 영향도 산정", "한국 개인정보 유형 사전 (주민등록번호 등)", "임성빈", "진행전", "2026-09-24", "2026-10-14"),
    ("메인 기능", "3단계 영향도 산정", "개인정보보호법 분류 매핑·가중치 설계", "임성빈", "진행전", "2026-10-08", "2026-10-28"),
    ("메인 기능", "3단계 영향도 산정", "개인정보보호위원회 처분 사례 수집·근거 문서화", "임성빈", "진행전", "2026-10-15", "2026-11-04"),
    ("메인 기능", "3단계 영향도 산정", "엔드포인트별 영향도 점수·수정 우선순위 산출", "임성빈", "진행전", "2026-10-29", "2026-11-18"),

    ("출력·통합", "리포트", "진단 리포트 (사람용) 제작", "공동", "진행전", "2026-11-05", "2026-11-18"),
    ("출력·통합", "리포트", "기계 판독용 출력 스키마 확정", "공동", "진행전", "2026-11-05", "2026-11-18"),
    ("출력·통합", "통합", "4단계 파이프라인 통합", "공동", "진행전", "2026-11-12", "2026-11-25"),
    ("출력·통합", "통합", "데모 시나리오 구성", "공동", "진행전", "2026-11-19", "2026-11-25"),

    ("평가·실험", "벤치마크", "crAPI 3건 + VAmPI 2건 벤치 고정", "양유상", "진행전", "2026-09-24", "2026-10-07"),
    ("평가·실험", "벤치마크", "평가 자동화 스크립트 제작", "양유상", "진행전", "2026-10-08", "2026-10-28"),
    ("평가·실험", "비교 실험", "ZAP Access Control·Autorize 입력량 비교", "양유상", "진행전", "2026-11-12", "2026-11-25"),
    ("평가·실험", "비교 실험", "필드 분류 베이스라인 Presidio 대조", "양유상", "진행전", "2026-11-12", "2026-11-25"),
    ("평가·실험", "비교 실험", "실험 결과 정리·수치화", "양유상", "진행전", "2026-11-26", "2026-12-02"),

    ("유지보수", "프로젝트 개선", "오탐·성능 분석 및 개선", "전체", "진행전", "2026-11-26", "2026-12-09"),
    ("유지보수", "프로젝트 개선", "고도화", "전체", "진행전", "2026-11-26", "2026-12-09"),

    ("기타", "문서화", "GitHub 저장소 개설·구조 정리", "심재훈", "진행전", "2026-09-24", "2026-09-30"),
    ("기타", "문서화", "프로젝트 문서화 (최종보고서 / README.md)", "전체", "진행전", "2026-12-03", "2026-12-09"),
    ("기타", "발표", "계획 발표자료 제작", "전체", "진행전", "2026-09-21", "2026-09-22"),
    ("기타", "발표", "계획 발표 (온라인 20:00)", "전체", "진행전", "2026-09-23", "2026-09-23"),
    ("기타", "발표", "중간 발표자료 제작", "전체", "진행전", "2026-10-22", "2026-10-28"),
    ("기타", "발표", "중간 발표", "미정", "진행전", "2026-10-29", "2026-10-29"),
    ("기타", "발표", "최종 발표자료 제작", "전체", "진행전", "2026-12-03", "2026-12-09"),
    ("기타", "발표", "최종 발표", "미정", "진행전", "2026-12-10", "2026-12-10"),
    ("기타", "보고서", "주간보고서 작성", "전체", "진행전", "2026-09-24", "2026-12-09"),
    ("기타", "WBS", "수정사항 생길 때마다 WBS 수정", "전체", "진행중", "2026-09-20", "2026-12-10"),
]

GRID_START = dt.date(2026, 9, 14)
GRID_END = dt.date(2026, 12, 12)
days = [GRID_START + dt.timedelta(days=i) for i in range((GRID_END - GRID_START).days + 1)]

SECTION_FILL = {
    "기획 및 설계": "A9D08E",
    "메인 기능": "FFD966",
    "출력·통합": "9DC3E6",
    "평가·실험": "E28E9C",
    "유지보수": "C9A0DC",
    "기타": "9DC3E6",
}
STATUS_FILL = {"완료": "C6EFCE", "진행중": "F8B26A", "진행전": "D9D9D9"}
BAR_FILL = PatternFill("solid", fgColor="FCE4A6")
HEAD_FILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="A6A6A6")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook()
ws = wb.active
ws.title = "WBS"

headers = ["구분", "작업", "기능 상세", "담당자", "완료여부", "기간"]
FIRST_DAY_COL = len(headers) + 1

for c, h in enumerate(headers, start=1):
    ws.cell(row=1, column=c, value=h)
    ws.merge_cells(start_row=1, start_column=c, end_row=2, end_column=c)

col = FIRST_DAY_COL
month_start = col
for i, d in enumerate(days):
    cell = ws.cell(row=2, column=col, value=d.day)
    if d.weekday() == 5:
        cell.font = Font(name=FONT, size=8, color="0070C0", bold=True)
    elif d.weekday() == 6:
        cell.font = Font(name=FONT, size=8, color="FF0000", bold=True)
    else:
        cell.font = Font(name=FONT, size=8)
    cell.alignment = Alignment(horizontal="center")
    cell.fill = HEAD_FILL
    cell.border = BORDER
    if i == len(days) - 1 or days[i + 1].month != d.month:
        ws.merge_cells(start_row=1, start_column=month_start, end_row=1, end_column=col)
        mc = ws.cell(row=1, column=month_start, value=str(d.month) + "월")
        mc.font = Font(name=FONT, size=9, bold=True)
        mc.alignment = Alignment(horizontal="center")
        mc.fill = HEAD_FILL
        month_start = col + 1
    col += 1

for c in range(1, FIRST_DAY_COL):
    cell = ws.cell(row=1, column=c)
    cell.font = Font(name=FONT, size=10, bold=True)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.fill = HEAD_FILL
    cell.border = BORDER

r = 3
for (sec, task, detail, owner, status, s, e) in rows:
    sd = dt.datetime.strptime(s, D).date()
    ed = dt.datetime.strptime(e, D).date()
    span = s[5:].replace("-", "/")
    if sd != ed:
        span = span + " ~ " + e[5:].replace("-", "/")
    vals = [sec, task, detail, owner, status, span]
    for c, v in enumerate(vals, start=1):
        cell = ws.cell(row=r, column=c, value=v)
        cell.font = Font(name=FONT, size=9, bold=(c == 1))
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=(c == 3))
        cell.border = BORDER
        if c == 1:
            cell.fill = PatternFill("solid", fgColor=SECTION_FILL[sec])
        if c == 5:
            cell.fill = PatternFill("solid", fgColor=STATUS_FILL[status])
    for i, d in enumerate(days):
        cell = ws.cell(row=r, column=FIRST_DAY_COL + i)
        cell.border = BORDER
        if sd <= d <= ed:
            cell.fill = BAR_FILL
    r += 1
LAST_ROW = r - 1


def merge_runs(colidx, keyfunc):
    start = 3
    for rr in range(3, LAST_ROW + 2):
        cur = keyfunc(rr) if rr <= LAST_ROW else None
        if cur != keyfunc(start):
            if rr - 1 > start:
                ws.merge_cells(start_row=start, start_column=colidx, end_row=rr - 1, end_column=colidx)
            start = rr


merge_runs(1, lambda rr: rows[rr - 3][0])
merge_runs(2, lambda rr: (rows[rr - 3][0], rows[rr - 3][1]))

lr = LAST_ROW + 2
notes = [
    "9/23 계획 발표까지만 확정 일정이다. 이후 날짜는 학기말 12월 중순 기준으로 역산한 제안이며, 중간·최종 발표 일자가 공지되면 고친다.",
    "담당자는 기획서 역할 분담표 기준이다. 심재훈 인가 탐지·LLM, 곽민경 트래픽·엔드포인트 복원, 임성빈 법 분류·영향도, 양유상 벤치마크·평가.",
    "완료여부는 완료 / 진행중 / 진행전 세 값만 쓴다. 기간 칸은 MM/DD ~ MM/DD 형식, 간트 막대는 기간 칸과 같은 날짜에 칠한다.",
]
ws.cell(row=lr, column=1, value="가정·규칙").font = Font(name=FONT, size=9, bold=True)
for i, t in enumerate(notes):
    ws.cell(row=lr + i, column=3, value=t).font = Font(name=FONT, size=9)

widths = {1: 14, 2: 22, 3: 46, 4: 10, 5: 10, 6: 16}
for c, w in widths.items():
    ws.column_dimensions[get_column_letter(c)].width = w
for i in range(len(days)):
    ws.column_dimensions[get_column_letter(FIRST_DAY_COL + i)].width = 2.6
ws.row_dimensions[1].height = 16
ws.row_dimensions[2].height = 16
ws.freeze_panes = ws.cell(row=3, column=FIRST_DAY_COL)
ws.sheet_view.zoomScale = 85

wb.save(OUT)
print("saved:", OUT, "task rows:", LAST_ROW - 2, "day columns:", len(days))
